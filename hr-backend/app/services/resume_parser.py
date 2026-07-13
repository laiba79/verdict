"""
Extracts plain text from an uploaded resume file (PDF or DOCX) so the rest
of the pipeline (LLM matching, keyword analysis) can work with plain text
regardless of the original format.

Deliberately does NOT try OCR for scanned/image-only PDFs -- that's a
different, heavier problem (needs tesseract or a vision model). We just
fail clearly and tell the uploader to try a text-based export instead.
"""

import io

import docx
from fastapi import HTTPException, UploadFile
from pypdf import PdfReader

MAX_RESUME_CHARS = 20_000  # generous cap; keeps LLM prompt size/cost bounded


async def extract_resume_text(file: UploadFile) -> str:
    filename = (file.filename or "").lower()
    content = await file.read()

    if not content:
        raise HTTPException(status_code=400, detail="The uploaded file is empty.")

    if filename.endswith(".pdf"):
        text = _extract_pdf(content)
    elif filename.endswith(".docx"):
        text = _extract_docx(content)
    elif filename.endswith(".doc"):
        raise HTTPException(
            status_code=400,
            detail="Legacy .doc files aren't supported. Please re-save as .docx or .pdf.",
        )
    else:
        raise HTTPException(
            status_code=400,
            detail="Unsupported file type. Please upload a PDF or DOCX resume.",
        )

    return text[:MAX_RESUME_CHARS]


def _extract_pdf(content: bytes) -> str:
    try:
        reader = PdfReader(io.BytesIO(content))
        text = "\n".join(page.extract_text() or "" for page in reader.pages)
    except Exception:
        raise HTTPException(
            status_code=400,
            detail="Couldn't read this PDF -- it may be corrupted or password-protected.",
        )

    text = text.strip()
    if not text:
        raise HTTPException(
            status_code=400,
            detail="No extractable text found in this PDF (it may be a scanned image without a text layer).",
        )
    return text


def _extract_docx(content: bytes) -> str:
    try:
        document = docx.Document(io.BytesIO(content))
        parts = [p.text for p in document.paragraphs]
        for table in document.tables:
            for row in table.rows:
                parts.extend(cell.text for cell in row.cells)
        text = "\n".join(p for p in parts if p.strip())
    except Exception:
        raise HTTPException(status_code=400, detail="Couldn't read this DOCX file -- it may be corrupted.")

    text = text.strip()
    if not text:
        raise HTTPException(status_code=400, detail="No extractable text found in this document.")
    return text
